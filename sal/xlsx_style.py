"""Shared openpyxl styling for SAL's formatted (non-raw-dump) reports.

sal/export.py's rows_to_csv/rows_to_xlsx are deliberately unstyled - a raw
data pull. sal/itgc_report.py and sal/rules_catalog.py both build styled,
multi-sheet reference documents instead, and want byte-identical visual
treatment (same header color, same wrap/border/column-width conventions)
so a reader doesn't see two different "SAL report" visual languages.
Extracted here once both had the same ~40 lines of styling duplicated.
"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

NAVY = "1F3864"
WHITE = "FFFFFF"

HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color=WHITE)
TITLE_FONT = Font(bold=True, size=14, color=NAVY)
SECTION_FONT = Font(bold=True, size=11, color=NAVY)
LABEL_FONT = Font(bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN_BORDER = Border(*(Side(style="thin", color="BFBFBF"),) * 4)


def style_header_row(ws: Worksheet, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = THIN_BORDER


def style_body(ws: Worksheet, first_row: int) -> None:
    for row in ws.iter_rows(min_row=first_row, max_row=ws.max_row):
        for cell in row:
            cell.alignment = WRAP
            cell.border = THIN_BORDER


def set_widths(ws: Worksheet, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
