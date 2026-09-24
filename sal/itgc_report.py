"""ITGC-formatted findings report (audit workpaper layout).

``sal/export.py``'s CSV/XLSX serializers are generic row dumps - fine for a
raw data pull, but not the shape an ITGC (IT General Controls) audit team
actually reviews. Two reference "SAP Initial Data Request" workbooks the
client's audit team shared (``Docs/*ITGC Data Request*.xlsx``) show what
they expect: a cover sheet naming client/application/audit period/
extraction method, an indexed control matrix (control objective -> test
performed -> population/exceptions -> reference sheet), a detailed
findings workpaper with disposition plus blank auditor-remarks columns for
their own use, and an extraction log evidencing how the underlying data
was pulled (their IDR literally asks for screenshots of RFC/GUI input
parameters - SAL's ``collection_runs`` ledger is the API-driven
equivalent of that evidence).

This module only formats data already fetched elsewhere - findings via
``sal.findings.list_findings()``, collection runs via the same inline
query ``sal/web/api.py`` already uses for ``/collection-runs/export`` -
it does no database access of its own, matching ``sal/export.py``'s
existing split between "fetch" (api.py) and "serialize" (here).

Scope, stated up front because an auditor will ask: SAL only continuously
monitors the SAP Security Audit Log (SM20), which evidences a single ITGC
control domain - Access to Programs and Data (logon/session monitoring,
sensitive transaction/table usage, and user-master change monitoring). It
has nothing to say about Program Change Management, Program Development,
or Computer Operations (backup/job scheduling); the cover sheet says so
explicitly rather than implying broader coverage than the tool has.
"""
import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from . import xlsx_style
from .export import sanitize_cell_value, sanitize_row
from .rules import RULES

# Maps each registered rule_key (sal/rules/registry.py's RULES) to the ITGC
# control objective it evidences. Kept here, not in registry.py, because
# this grouping is a reporting concern (how an auditor should read the
# rule catalog) rather than a detection concern.
CONTROL_OBJECTIVES = [
    {
        "ref": "1.1",
        "title": "Logon & Session Access Monitoring",
        "description": (
            "Access is monitored for unusual logon timing/frequency, "
            "unrecognized source terminals/IPs, repeated authentication "
            "failures, and multiple users sharing one source."
        ),
        "rules": ["login_time", "login_frequency", "new_source", "login_attack", "shared_ip"],
    },
    {
        "ref": "1.2",
        "title": "Segregation of Duties & Sensitive Transaction Usage",
        "description": (
            "Transaction execution is monitored for activity outside a "
            "user's assigned role/baseline pattern, first-time use of a "
            "sensitive transaction, and abnormal activity volume."
        ),
        "rules": [
            "unusual_transaction", "first_time_sensitive_transaction", "activity_volume",
            "out_of_context_transaction", "out_of_context_no_role_data",
            "critical_transaction_usage",
        ],
    },
    {
        "ref": "1.3",
        "title": "Sensitive Data Access & Export Monitoring",
        "description": "Access to sensitive tables and data extraction/export activity is monitored.",
        "rules": ["sensitive_table_access", "data_export"],
    },
    {
        "ref": "1.4",
        "title": "Privileged / User Master Data Change Monitoring",
        "description": ("Bulk or unusual changes to user master records are monitored for "
                         "unauthorized provisioning."),
        "rules": ["mass_user_changes"],
    },
]

UNMAPPED_REF = "1.9"
UNMAPPED_TITLE = "Other / Unmapped Detection Rule"

DISPOSITION_LEGEND = [
    ("open", "Not yet reviewed by an analyst."),
    ("true_positive", "Reviewed and confirmed as a genuine exception."),
    ("false_positive", "Reviewed and determined not to be an exception."),
    ("whitelisted", "Matches a time-bound, analyst-approved whitelist rule."),
]

SEVERITY_LEGEND = [
    ("High", ("Warrants prompt analyst review - e.g. privileged access, mass changes, "
              "sensitive data movement.")),
    ("Medium", "Notable deviation from baseline behavior; reviewed during normal triage."),
    ("Low", "Informational deviation; reviewed at lower priority."),
]

def _rule_control_lookup() -> dict[str, dict]:
    lookup = {}
    for objective in CONTROL_OBJECTIVES:
        for rule_key in objective["rules"]:
            lookup[rule_key] = objective
    return lookup


def _rule_label(rule_key: str) -> str:
    for key, label, _fn in RULES:
        if key == rule_key:
            return label
    return rule_key


def _format_evidence(evidence: dict | None) -> str:
    if not evidence:
        return ""
    parts = []
    for key in sorted(evidence):
        value = evidence[key]
        if isinstance(value, float):
            value = round(value, 2)
        parts.append(f"{key}: {value}")
    return "; ".join(parts)


def _disposition_rationale(finding: dict) -> str:
    parts = [p for p in (finding.get("reason_code"), finding.get("analyst_notes")) if p]
    return " — ".join(parts)


def _min_max(values: list[str]) -> tuple[str, str]:
    clean = sorted(v for v in values if v)
    if not clean:
        return "", ""
    return clean[0], clean[-1]


def _format_sap_date(value: str) -> str:
    """dat_from/dat_to are SAP DATS strings (YYYYMMDD, per api.py's own
    _DATE_RE validation) - render as YYYY-MM-DD for an audit-report
    reader. Falls back to the raw value for anything else rather than
    raising, since this only feeds a display string.
    """
    if value and len(value) == 8 and value.isdigit():
        return f"{value[0:4]}-{value[4:6]}-{value[6:8]}"
    return value


def _build_cover_sheet(wb: Workbook, *, org_name: str, application: str,
                        system_id: str | None, client: str | None,
                        environment: str | None, collection_runs: list[dict],
                        findings: list[dict], generated_by: str, generated_at: str) -> None:
    ws = wb.active
    ws.title = "Cover Sheet"
    xlsx_style.set_widths(ws, [3, 30, 80])

    row = 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    title_cell = ws.cell(row=row, column=1,
                          value="SAP Security Audit Log (SM20) — ITGC Access Monitoring Report")
    title_cell.font = xlsx_style.TITLE_FONT
    row += 2

    _, period_to = _min_max([r.get("dat_to") for r in collection_runs if r.get("status") == "success"])
    period_from, _ = _min_max([r.get("dat_from") for r in collection_runs if r.get("status") == "success"])
    exceptions_from, exceptions_to = _min_max([f.get("detected_at") for f in findings])

    system_label = system_id or "All configured systems"
    if system_id and environment:
        system_label = f"{system_id} ({environment})"

    kv_rows = [
        ("Client / Organization", org_name),
        ("Application", application),
        ("System", system_label),
        ("Client (SAP mandt)", client or "All"),
        ("Testing Period (SM20 data collected)",
         f"{_format_sap_date(period_from)} to {_format_sap_date(period_to)}" if period_from
         else "No successful collection runs in scope"),
        ("Exceptions Observed (event date range)",
         f"{exceptions_from} to {exceptions_to}" if exceptions_from else "No findings in scope"),
        ("Report Generated By", generated_by),
        ("Report Generated On (UTC)", generated_at),
        ("Data Source", "SAP Security Audit Log (transaction SM20), extracted read-only via RFC "
                         "function module RSAU_API_GET_LOG_DATA"),
        ("Detection Methodology", "Deterministic, rule-based detection (no ML/LLM) - see the "
                                  "'Control Matrix' sheet for the control objective each rule evidences"),
        ("Reporting Tool", "Security Audit Log (SAL) Tool"),
    ]
    for label, value in kv_rows:
        ws.cell(row=row, column=2, value=label).font = xlsx_style.LABEL_FONT
        ws.cell(row=row, column=3, value=sanitize_cell_value(value)).alignment = xlsx_style.WRAP
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="Scope & Limitations").font = xlsx_style.SECTION_FONT
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    scope_cell = ws.cell(row=row, column=1, value=(
        "This report covers continuous monitoring of the SAP Security Audit Log (SM20) only. "
        "It evidences the ITGC control objective of monitoring access to programs and data "
        "(logon activity, sensitive transaction/table usage, and user master data changes) and "
        "does not cover Program Change Management, Program Development, or Computer Operations "
        "(backup/job scheduling) ITGC domains. All findings are generated by deterministic rules "
        "against retained SM20 event history; a finding reflects an analyst's judgment only once "
        "it has been reviewed and dispositioned (see 'Disposition' / 'Reviewed By' / 'Reviewed On' "
        "in the Findings Detail sheet) - an 'open' status means analyst review is still pending, "
        "not that the condition has been ruled out. This tool is read-only and never writes back "
        "to the source SAP system."
    ))
    scope_cell.alignment = xlsx_style.WRAP
    ws.row_dimensions[row].height = 90
    row += 2

    ws.cell(row=row, column=1, value="Severity Ratings").font = xlsx_style.SECTION_FONT
    row += 1
    for label, desc in SEVERITY_LEGEND:
        ws.cell(row=row, column=2, value=label).font = xlsx_style.LABEL_FONT
        ws.cell(row=row, column=3, value=desc).alignment = xlsx_style.WRAP
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="Disposition Statuses").font = xlsx_style.SECTION_FONT
    row += 1
    for label, desc in DISPOSITION_LEGEND:
        ws.cell(row=row, column=2, value=label).font = xlsx_style.LABEL_FONT
        ws.cell(row=row, column=3, value=desc).alignment = xlsx_style.WRAP
        row += 1


def _build_control_matrix(wb: Workbook, findings: list[dict]) -> None:
    ws = wb.create_sheet("Control Matrix")
    headers = [
        "S.No", "Control Ref", "ITGC Control Objective", "Control Description",
        "Detection Rules Applied", "Population (Total Findings)",
        "Open (Pending Review)", "Confirmed Exception", "Cleared (False Positive)",
        "Whitelisted (Pre-Approved)", "Reference Sheet",
    ]
    ws.append(headers)
    xlsx_style.style_header_row(ws, 1, len(headers))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    counts_by_rule: dict[str, dict[str, int]] = {}
    for f in findings:
        bucket = counts_by_rule.setdefault(
            f["rule_key"], {"open": 0, "true_positive": 0, "false_positive": 0, "whitelisted": 0}
        )
        if f.get("status") in bucket:
            bucket[f["status"]] += 1

    for i, objective in enumerate(CONTROL_OBJECTIVES, start=1):
        totals = {"open": 0, "true_positive": 0, "false_positive": 0, "whitelisted": 0}
        labels = []
        for rule_key in objective["rules"]:
            labels.append(_rule_label(rule_key))
            rule_counts = counts_by_rule.get(rule_key, {})
            for status in totals:
                totals[status] += rule_counts.get(status, 0)
        population = sum(totals.values())
        ws.append(sanitize_row([
            i, objective["ref"], objective["title"], objective["description"],
            ", ".join(labels), population,
            totals["open"], totals["true_positive"], totals["false_positive"], totals["whitelisted"],
            "Findings Detail",
        ]))
    xlsx_style.set_widths(ws, [6, 10, 32, 48, 42, 14, 12, 12, 12, 14, 16])
    xlsx_style.style_body(ws, first_row=2)


def _build_findings_detail(wb: Workbook, findings: list[dict], control_lookup: dict[str, dict]) -> None:
    ws = wb.create_sheet("Findings Detail")
    headers = [
        "S.No", "Control Ref", "ITGC Control Objective", "Test Performed (Rule)",
        "Risk Rating", "System", "Client", "SAP User ID",
        "Event Date/Time (UTC)", "First Observed", "Last Observed",
        "Observation / Exception Description", "Supporting Evidence",
        "Disposition", "Disposition Rationale", "Reviewed By", "Reviewed On",
        "Auditor Conclusion", "Auditor Remarks",
    ]
    ws.append(headers)
    xlsx_style.style_header_row(ws, 1, len(headers))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    for i, f in enumerate(findings, start=1):
        objective = control_lookup.get(f["rule_key"], {"ref": UNMAPPED_REF, "title": UNMAPPED_TITLE})
        ws.append(sanitize_row([
            i, objective["ref"], objective["title"], f.get("rule_label", f["rule_key"]),
            f.get("severity"), f.get("source_system"), f.get("client"), f.get("user_id"),
            f.get("detected_at"), f.get("first_seen_at"), f.get("last_seen_at"),
            f.get("summary"), _format_evidence(f.get("evidence")),
            f.get("status"), _disposition_rationale(f), f.get("disposed_by"), f.get("disposed_at"),
            "", "",
        ]))
    xlsx_style.set_widths(ws, [6, 10, 30, 30, 10, 10, 8, 14, 20, 20, 20, 45, 45, 14, 28, 14, 20, 20, 30])
    xlsx_style.style_body(ws, first_row=2)


def _build_extraction_log(wb: Workbook, collection_runs: list[dict]) -> None:
    ws = wb.create_sheet("Extraction Log")
    headers = [
        "Run ID", "System", "Client", "Period From", "Period To",
        "Started At", "Finished At", "Status", "Rows Collected",
        "Run By", "Extraction Parameters", "Error",
    ]
    ws.append(headers)
    xlsx_style.style_header_row(ws, 1, len(headers))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    for r in collection_runs:
        ws.append(sanitize_row([
            r.get("run_id"), r.get("source_system"), r.get("client"),
            _format_sap_date(r.get("dat_from") or ""), _format_sap_date(r.get("dat_to") or ""),
            r.get("started_at"), r.get("finished_at"),
            r.get("status"), r.get("row_count"), r.get("actor"),
            r.get("filters_json") or "", r.get("error_message") or "",
        ]))
    xlsx_style.set_widths(ws, [22, 10, 8, 14, 14, 20, 20, 12, 12, 14, 40, 30])
    xlsx_style.style_body(ws, first_row=2)


def build_itgc_workbook(
    findings: list[dict],
    collection_runs: list[dict],
    *,
    system_id: str | None,
    client: str | None,
    environment: str | None = None,
    org_name: str = "Bristlecone Group",
    application: str = "SAP ECC 6.0",
    generated_by: str = "unknown",
    generated_at: str | None = None,
) -> bytes:
    """Render findings + collection-run history into the ITGC workpaper
    layout: Cover Sheet, Control Matrix, Findings Detail, Extraction Log.
    """
    generated_at = generated_at or datetime.now(timezone.utc).isoformat()
    wb = Workbook()
    control_lookup = _rule_control_lookup()

    _build_cover_sheet(
        wb, org_name=org_name, application=application, system_id=system_id,
        client=client, environment=environment, collection_runs=collection_runs,
        findings=findings, generated_by=generated_by, generated_at=generated_at,
    )
    _build_control_matrix(wb, findings)
    _build_findings_detail(wb, findings, control_lookup)
    _build_extraction_log(wb, collection_runs)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
