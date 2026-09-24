"""Detection rules catalog (plain-English reference document).

Findings tell an analyst *what* fired; nothing in the running app explains
*why* a rule exists or how its detection logic works short of reading the
Python source in sal/rules/ or the prose in Docs/Rules.md. This module
turns that into a downloadable, styled Excel workbook so an analyst (or an
auditor asking "how does this tool decide something is unusual?") has a
single self-contained reference document, without needing repo access.

RULE_DETAILS's descriptions are written from the actual detection logic in
each sal/rules/*.py module (read directly, not copied from Docs/Rules.md's
higher-level table) so they describe what the code does, not what the
original use-case brief intended.
"""
import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from . import xlsx_style
from .export import sanitize_cell_value, sanitize_row
from .rules import RULES

# Keyed by rule_key (sal/rules/registry.py's RULES) - see that module's
# docstring for the registry itself. use_case is the numbering from the
# original Docs/Security_AI_UseCases_v1.xlsx blueprint (Docs/Rules.md keeps
# the two in sync); source_file is repo-relative for a reader who wants to
# go straight to the code.
RULE_DETAILS = {
    "login_attack": {
        "use_case": "#9",
        "what_it_detects": (
            "Repeated failed logon attempts for a user followed by a successful logon within "
            "a short window (potential brute-force/credential-guessing), or SAP's own "
            "account-lockout event after too many failed password checks."
        ),
        "how_it_works": (
            "Watches AU1 (logon success), AU2 (logon failed), and AUM (account locked) events "
            "per user. Raises 'account_locked' whenever SAP itself locks the account, and "
            "'potential_login_attack' whenever 3+ failed logons occur within a 15-minute "
            "window immediately before a successful one."
        ),
        "sm19_signal": "AU1, AU2, AUM (Login)",
        "typical_severity": "High",
        "notes": "",
        "source_file": "sal/rules/login_attack.py",
    },
    "mass_user_changes": {
        "use_case": "#12",
        "what_it_detects": (
            "A burst of user-master-record changes (user creation, authorization changes, "
            "unlocks, password resets) happening close together in time - consistent with "
            "bulk or unauthorized user provisioning."
        ),
        "how_it_works": (
            "Groups all 'User master changes'-class events into incidents (events less than "
            "15 minutes apart are merged into one incident); raises one finding per incident "
            "once it reaches 5+ events, rather than one finding per event in a sustained burst."
        ),
        "sm19_signal": "User master changes (AUD/AU7/AUB/AUA/BU2 confirmed in real data)",
        "typical_severity": "Medium",
        "notes": ("v1 counts overall event volume, not distinct target usernames - the target "
                  "username isn't in a dedicated column for these codes yet."),
        "source_file": "sal/rules/mass_user_changes.py",
    },
    "new_source": {
        "use_case": "#3",
        "what_it_detects": "A user logging on from an IP address never seen for that user before.",
        "how_it_works": (
            "The first 3 logons (with an IP) for each user establish their known-IP baseline; "
            "any later logon from an unlisted IP is flagged. Blank-IP background/batch logons "
            "are ignored."
        ),
        "sm19_signal": "AU1 (Login)",
        "typical_severity": "Medium",
        "notes": "",
        "source_file": "sal/rules/new_source.py",
    },
    "unusual_transaction": {
        "use_case": "#4",
        "what_it_detects": (
            "A user running a transaction code they've never run before, after an initial "
            "learning period."
        ),
        "how_it_works": (
            "The first 5 distinct tcodes a user runs establish their baseline; any later "
            "never-before-seen tcode is flagged."
        ),
        "sm19_signal": "AU3 (Transaction start)",
        "typical_severity": "Medium",
        "notes": "",
        "source_file": "sal/rules/first_time_transaction.py",
    },
    "first_time_sensitive_transaction": {
        "use_case": "#11",
        "what_it_detects": (
            "The very first time a user runs a transaction on SAL's critical-transaction "
            "catalogue (e.g. SU01, SE38) - no learning period; the first occurrence is itself "
            "the finding."
        ),
        "how_it_works": (
            "Same engine as 'Unusual transaction', restricted to the critical-transaction "
            "list, with the baseline/learning-period disabled."
        ),
        "sm19_signal": "AU3 (Transaction start)",
        "typical_severity": "High",
        "notes": "",
        "source_file": "sal/rules/first_time_transaction.py",
    },
    "shared_ip": {
        "use_case": "#7",
        "what_it_detects": (
            "Multiple distinct users logging on from the same source IP within a short window "
            "- consistent with a shared jump host/proxy or credential abuse from one origin."
        ),
        "how_it_works": (
            "Groups successful logons by IP; flags an IP once 3+ distinct users have logged on "
            "from it within a rolling 10-minute window. The SAP application server's own IP "
            "and loopback addresses are excluded."
        ),
        "sm19_signal": "AU1 (Login)",
        "typical_severity": "Medium",
        "notes": "",
        "source_file": "sal/rules/shared_ip_multi_user.py",
    },
    "sensitive_table_access": {
        "use_case": "#8",
        "what_it_detects": "A user accessing (display/change/delete) a table on SAL's sensitive-tables catalogue.",
        "how_it_works": (
            "Filters generic table-access events (DU9) to those naming a sensitive table; "
            "severity is High for change/delete activity, Low for a plain display."
        ),
        "sm19_signal": "DU9 (Other events)",
        "typical_severity": "Low or High (activity-dependent)",
        "notes": "",
        "source_file": "sal/rules/export.py",
    },
    "data_export": {
        "use_case": "#8",
        "what_it_detects": (
            "A user downloading data from SAP, especially when it follows a sensitive-table "
            "access shortly before."
        ),
        "how_it_works": (
            "Correlates download events (AUY) against prior sensitive-table access (DU9) by "
            "the same user within 15 minutes; raises 'sensitive_data_export' (High) when "
            "correlated, or the lower-severity 'data_export' (Low) for an uncorrelated "
            "download that's still worth a record."
        ),
        "sm19_signal": "AUY, DU9 (Other events)",
        "typical_severity": "Low or High (correlation-dependent)",
        "notes": "",
        "source_file": "sal/rules/export.py",
    },
    "login_time": {
        "use_case": "#1",
        "what_it_detects": "A user logging on at a time of day that's a significant departure from their own normal pattern.",
        "how_it_works": (
            "Tracks each user's login hour-of-day as a running mean/standard deviation; once a "
            "user has 5+ prior logons, a new logon more than 2.5 standard deviations from "
            "their usual time is flagged."
        ),
        "sm19_signal": "AU1 (Login)",
        "typical_severity": "Medium",
        "notes": "",
        "source_file": "sal/rules/login_time.py",
    },
    "login_frequency": {
        "use_case": "#2",
        "what_it_detects": "A day on which a user logs on far more often than their own historical daily average.",
        "how_it_works": (
            "Tracks each user's daily logon (AU1) count as a running mean/standard deviation; "
            "flags a day 2+ standard deviations above that user's own average."
        ),
        "sm19_signal": "AU1 (Login)",
        "typical_severity": "Medium",
        "notes": "",
        "source_file": "sal/rules/daily_count_baseline.py",
    },
    "activity_volume": {
        "use_case": "#5",
        "what_it_detects": "A day on which a user's overall transaction activity is far above their own historical daily average.",
        "how_it_works": "Same running-baseline engine as 'Unusual login frequency', counting daily transaction-start (AU3) events per user.",
        "sm19_signal": "AU3 (Transaction start)",
        "typical_severity": "Medium",
        "notes": "",
        "source_file": "sal/rules/daily_count_baseline.py",
    },
    "out_of_context_transaction": {
        "use_case": "#14",
        "what_it_detects": "A user running a critical transaction that isn't covered by any SAP role currently assigned to them.",
        "how_it_works": (
            "Cross-references every critical-transaction event against the user's "
            "currently-assigned role/tcode data (refreshed daily); flags when the user has "
            "roles on record but none include that tcode. Checks current access, not access "
            "at the time the event happened, so an access change since surfaces new findings "
            "on retained history."
        ),
        "sm19_signal": "AU3 (Transaction start) + PFCG role/tcode data (AGR_USERS/AGR_TCODES)",
        "typical_severity": "High",
        "notes": ("Wording is deliberately conservative ('not currently assigned via any "
                  "role', never 'unauthorized') - a tcode in a role's menu isn't the same as "
                  "full authorization-object-level access."),
        "source_file": "sal/rules/out_of_context_transaction.py",
    },
    "out_of_context_no_role_data": {
        "use_case": "#14",
        "what_it_detects": (
            "A user ran a critical transaction but has zero SAP role assignments on record at "
            "all - a rarer, more unusual condition than 'Out-of-context transaction usage' "
            "(roles exist but don't cover it)."
        ),
        "how_it_works": "Same cross-reference as 'Out-of-context transaction usage', kept as a distinct rule so this rarer case isn't hidden inside the more common one.",
        "sm19_signal": "AU3 (Transaction start) + PFCG role/tcode data",
        "typical_severity": "High",
        "notes": "",
        "source_file": "sal/rules/out_of_context_transaction.py",
    },
    "critical_transaction_usage": {
        "use_case": "N/A - added post-blueprint (2026-09-15)",
        "what_it_detects": (
            "Every execution of a transaction on SAL's critical-transaction catalogue, by "
            "anyone, unconditionally - authorized or not, first occurrence or five-hundredth."
        ),
        "how_it_works": (
            "Flags every critical-transaction AU3 event with no first-time or role/profile "
            "check at all - deliberately unlike 'First-time sensitive transaction' (fires once "
            "per user+tcode, ever) and 'Out-of-context transaction usage' (fires only when not "
            "covered by a current role). Requested directly for a comprehensive audit trail of "
            "privileged-transaction usage rather than an anomaly signal."
        ),
        "sm19_signal": "AU3 (Transaction start)",
        "typical_severity": "High",
        "notes": ("Fires repeatedly for routine, fully-authorized admin activity by design - "
                  "that volume is intentional, not a defect. Use the whitelist for a known, "
                  "accepted pattern that shouldn't keep generating new open findings."),
        "source_file": "sal/rules/critical_transaction_usage.py",
    },
}

NOT_YET_IMPLEMENTED = [
    ("#6", "Deviation from peer users",
     "Needs a peer-group baseline built on top of the role data #14 already ingests."),
    ("#10", "Firefighter/GRC session analysis",
     "Needs GRC EAM tables, which this tool does not ingest. Deferred pending separate approval."),
    ("#13", "Natural-language \"ask the tool\" query",
     "An LLM-layer feature, deliberately last since every current finding is already "
     "explainable via deterministic templates without one."),
]


def _rule_detail(rule_key: str, label: str) -> dict:
    return RULE_DETAILS.get(rule_key, {
        "use_case": "",
        "what_it_detects": f"(no catalog entry yet for '{label}')",
        "how_it_works": "",
        "sm19_signal": "",
        "typical_severity": "",
        "notes": "",
        "source_file": "",
    })


def _build_overview_sheet(wb: Workbook, generated_at: str) -> None:
    ws = wb.active
    ws.title = "Overview"
    xlsx_style.set_widths(ws, [3, 22, 90])

    row = 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    ws.cell(row=row, column=1, value="SAL Detection Rules Catalog").font = xlsx_style.TITLE_FONT
    row += 2

    kv_rows = [
        ("Report Generated On (UTC)", generated_at),
        ("Total Active Rules", str(len(RULES))),
        ("Detection Methodology", "Deterministic, rule-based detection only - no ML/LLM. "
                                   "Every finding's 'why flagged' text is a plain Python "
                                   "string template, not a model output."),
        ("Data Source", "SAP Security Audit Log (transaction SM20), extracted read-only via "
                         "RFC function module RSAU_API_GET_LOG_DATA"),
    ]
    for label, value in kv_rows:
        ws.cell(row=row, column=2, value=label).font = xlsx_style.LABEL_FONT
        ws.cell(row=row, column=3, value=sanitize_cell_value(value)).alignment = xlsx_style.WRAP
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="Not Yet Implemented").font = xlsx_style.SECTION_FONT
    row += 1
    for use_case, title, reason in NOT_YET_IMPLEMENTED:
        ws.cell(row=row, column=2, value=sanitize_cell_value(f"{use_case} {title}")).font = xlsx_style.LABEL_FONT
        ws.cell(row=row, column=3, value=sanitize_cell_value(reason)).alignment = xlsx_style.WRAP
        row += 1


def _build_rules_sheet(wb: Workbook) -> None:
    ws = wb.create_sheet("Detection Rules")
    headers = [
        "S.No", "Use Case #", "Rule Key", "Rule Label", "What It Detects",
        "How It Works", "SM19 Signal(s) Used", "Typical Severity", "Notes / Caveats",
        "Source File",
    ]
    ws.append(headers)
    xlsx_style.style_header_row(ws, 1, len(headers))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"

    for i, (rule_key, label, _fn) in enumerate(RULES, start=1):
        detail = _rule_detail(rule_key, label)
        ws.append(sanitize_row([
            i, detail["use_case"], rule_key, label,
            detail["what_it_detects"], detail["how_it_works"],
            detail["sm19_signal"], detail["typical_severity"],
            detail["notes"], detail["source_file"],
        ]))
    xlsx_style.set_widths(ws, [6, 10, 34, 34, 55, 55, 34, 22, 45, 34])
    xlsx_style.style_body(ws, first_row=2)


def build_rules_catalog_workbook(generated_at: str | None = None) -> bytes:
    """Render the active RULES registry into a plain-English reference
    workbook: Overview (methodology + not-yet-implemented use cases) and
    Detection Rules (one row per registered rule).
    """
    generated_at = generated_at or datetime.now(timezone.utc).isoformat()
    wb = Workbook()
    _build_overview_sheet(wb, generated_at)
    _build_rules_sheet(wb)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
