"""Regression tests for sal/export.py (Blueprint v2 Phase C).

CSV/XLSX serialization is pure and doesn't touch the database, so these
tests use plain in-memory row dicts rather than the isolated_db fixture's
throwaway sqlite file.
"""
import csv
import io
import json

from openpyxl import load_workbook

import sal.export as export


def test_rows_to_csv_writes_header_and_rows():
    headers = ["a", "b"]
    rows = [{"a": "1", "b": "2"}, {"a": "3", "b": "4"}]

    text = export.rows_to_csv(headers, rows)
    parsed = list(csv.DictReader(io.StringIO(text)))

    assert parsed == rows


def test_rows_to_csv_ignores_extra_fields_not_in_headers():
    headers = ["a"]
    rows = [{"a": "1", "unexpected": "should be dropped"}]

    text = export.rows_to_csv(headers, rows)
    parsed = list(csv.DictReader(io.StringIO(text)))

    assert parsed == [{"a": "1"}]


def test_rows_to_xlsx_writes_header_and_rows():
    headers = ["a", "b"]
    rows = [{"a": 1, "b": "x"}, {"a": 2, "b": "y"}]

    data = export.rows_to_xlsx(headers, rows, sheet_title="MySheet")
    wb = load_workbook(io.BytesIO(data))
    ws = wb["MySheet"]

    assert [c.value for c in ws[1]] == headers
    assert [c.value for c in ws[2]] == [1, "x"]
    assert [c.value for c in ws[3]] == [2, "y"]


def test_flatten_finding_serializes_evidence_dict_to_json_string():
    finding = {
        "finding_key": "abc", "rule_key": "new_source", "evidence": {"ip": "10.0.0.1", "count": 3},
    }

    flattened = export.flatten_finding(finding)

    assert isinstance(flattened["evidence"], str)
    assert json.loads(flattened["evidence"]) == {"ip": "10.0.0.1", "count": 3}
    assert flattened["finding_key"] == "abc"


def test_flatten_finding_handles_missing_evidence():
    flattened = export.flatten_finding({"finding_key": "abc"})
    assert flattened["evidence"] == "{}"


def test_sanitize_cell_value_defangs_formula_trigger_prefixes():
    for dangerous in ("=cmd", "+cmd", "-cmd", "@cmd", "\tcmd", "\rcmd"):
        assert export.sanitize_cell_value(dangerous) == "'" + dangerous


def test_sanitize_cell_value_leaves_ordinary_values_untouched():
    assert export.sanitize_cell_value("ordinary text") == "ordinary text"
    assert export.sanitize_cell_value(42) == 42
    assert export.sanitize_cell_value(None) is None


def test_rows_to_csv_defangs_formula_injection():
    headers = ["a"]
    rows = [{"a": "=HYPERLINK(\"http://evil\")"}]
    text = export.rows_to_csv(headers, rows)
    parsed = list(csv.DictReader(io.StringIO(text)))
    assert parsed[0]["a"] == "'=HYPERLINK(\"http://evil\")"


def test_rows_to_xlsx_defangs_formula_injection():
    headers = ["a"]
    rows = [{"a": "=HYPERLINK(\"http://evil\")"}]
    data = export.rows_to_xlsx(headers, rows)
    wb = load_workbook(io.BytesIO(data))
    ws = wb.active
    assert ws["A2"].value == "'=HYPERLINK(\"http://evil\")"
