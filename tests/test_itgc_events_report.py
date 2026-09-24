"""Regression tests for sal/itgc_events_report.py.

Formatting is pure and doesn't touch the database, so these tests use
synthetic event dicts rather than the isolated_db fixture's sqlite file.
"""
import io

from openpyxl import load_workbook

import sal.itgc_events_report as itgc_events_report
from sal.export import EVENTS_HEADERS


def _event(event_timestamp="2026-01-05 08:00:00", source_system="S23", client="100",
           user_id="ALICE", transaction_code="SU01", message="ok"):
    return {
        "event_timestamp": event_timestamp, "source_system": source_system, "client": client,
        "user_id": user_id, "user_email": None, "transaction_code": transaction_code,
        "msg_code": "AU1", "area": "AU", "event_class": "Dialog Logon", "severity": "H",
        "program": None, "terminal": "TERM1", "ip_address": "10.0.0.1", "message": message,
        "param1": None, "param2": None, "param3": None, "instance": "X",
        "log_tstmp": "20260105080000.0", "counter": 0, "collection_run_id": "run-1",
        "inserted_at": "2026-01-05T08:00:01+00:00",
    }


def test_build_events_workbook_produces_expected_sheets():
    data = itgc_events_report.build_events_workbook(
        [_event()], system_id="S23", client="100", dat_from="20260101", dat_to="20260131",
        filters={"transaction_code": ["SU01"]}, environment="Sandbox",
        generated_by="tester", generated_at="2026-09-12T00:00:00+00:00",
    )
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["Cover Sheet", "Event Log"]

    log = wb["Event Log"]
    header = [c.value for c in log[1]]
    assert header == EVENTS_HEADERS
    assert log.max_row == 2


def test_build_events_workbook_handles_no_events():
    data = itgc_events_report.build_events_workbook(
        [], system_id="S23", client="100", dat_from="20260101", dat_to="20260101", filters={},
    )
    wb = load_workbook(io.BytesIO(data))
    assert wb["Event Log"].max_row == 1  # header only


def test_cover_sheet_shows_extraction_parameters():
    data = itgc_events_report.build_events_workbook(
        [_event()], system_id="S23", client="100", dat_from="20260101", dat_to="20260131",
        filters={"transaction_code": ["SU01"], "user": ["ALICE", "BOB"]},
    )
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Cover Sheet"]
    row = next(r for r in ws.iter_rows(values_only=True) if r[1] == "Extraction Parameters")
    assert "SU01" in row[2]
    assert "ALICE" in row[2] and "BOB" in row[2]


def test_cover_sheet_reports_no_filters_when_none_given():
    data = itgc_events_report.build_events_workbook(
        [], system_id="S23", client=None, dat_from="20260101", dat_to="20260101", filters={},
    )
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Cover Sheet"]
    row = next(r for r in ws.iter_rows(values_only=True) if r[1] == "Extraction Parameters")
    assert "None" in row[2]


def test_format_sap_date_converts_yyyymmdd_to_iso():
    assert itgc_events_report._format_sap_date("20260601") == "2026-06-01"
    assert itgc_events_report._format_sap_date("") == ""
    assert itgc_events_report._format_sap_date(None) == ""


def test_defangs_formula_injection_in_event_message():
    data = itgc_events_report.build_events_workbook(
        [_event(message="=HYPERLINK(\"http://evil\")")],
        system_id="S23", client="100", dat_from="20260101", dat_to="20260101", filters={},
    )
    wb = load_workbook(io.BytesIO(data))
    log = wb["Event Log"]
    message_col = EVENTS_HEADERS.index("message")
    row = [c.value for c in log[2]]
    assert row[message_col].startswith("'=")
