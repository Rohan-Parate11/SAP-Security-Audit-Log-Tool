"""Regression tests for sal/rules_catalog.py.

Formatting is pure and doesn't touch the database, so no isolated_db
fixture is needed here.
"""
import io

from openpyxl import load_workbook

import sal.rules_catalog as rules_catalog
from sal.rules import RULES


def test_rule_details_covers_every_registered_rule():
    """Drift guard, mirroring sal/itgc_report.py's control-map drift test -
    a rule added to the registry without a RULE_DETAILS entry would
    otherwise silently render as '(no catalog entry yet for this rule)'.
    """
    registered_keys = {key for key, _label, _fn in RULES}
    assert registered_keys <= set(rules_catalog.RULE_DETAILS)


def test_build_rules_catalog_workbook_produces_expected_sheets():
    data = rules_catalog.build_rules_catalog_workbook(generated_at="2026-09-11T00:00:00+00:00")
    wb = load_workbook(io.BytesIO(data))

    assert wb.sheetnames == ["Overview", "Detection Rules"]

    detail = wb["Detection Rules"]
    assert detail.max_row == 1 + len(RULES)
    header = [c.value for c in detail[1]]
    assert header == [
        "S.No", "Use Case #", "Rule Key", "Rule Label", "What It Detects",
        "How It Works", "SM19 Signal(s) Used", "Typical Severity", "Notes / Caveats",
        "Source File",
    ]


def test_every_rule_row_has_a_non_placeholder_description():
    data = rules_catalog.build_rules_catalog_workbook()
    wb = load_workbook(io.BytesIO(data))
    detail = wb["Detection Rules"]

    for row in detail.iter_rows(min_row=2, values_only=True):
        rule_key, what_it_detects = row[2], row[4]
        assert what_it_detects != "(no catalog entry yet for this rule)", rule_key
        assert what_it_detects  # non-empty


def test_overview_sheet_lists_not_yet_implemented_use_cases():
    data = rules_catalog.build_rules_catalog_workbook()
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Overview"]

    labels = [row[1] for row in ws.iter_rows(min_row=1, max_col=2, values_only=True) if row[1]]
    assert any("#6" in label for label in labels)
    assert any("#10" in label for label in labels)
    assert any("#13" in label for label in labels)


def test_defangs_formula_injection_in_notes():
    """RULE_DETAILS values are static/code-controlled, but this workbook is
    still built with the same shared sanitize_cell_value helper as every
    other SAL export for consistency - this pins that it's actually wired
    up, not just present in the file.
    """
    original = dict(rules_catalog.RULE_DETAILS["new_source"])
    rules_catalog.RULE_DETAILS["new_source"]["notes"] = "=cmd|'/bin/calc'!A1"
    try:
        data = rules_catalog.build_rules_catalog_workbook()
    finally:
        rules_catalog.RULE_DETAILS["new_source"] = original

    wb = load_workbook(io.BytesIO(data))
    detail = wb["Detection Rules"]
    row = next(r for r in detail.iter_rows(min_row=2, values_only=True) if r[2] == "new_source")
    assert row[8].startswith("'=")
