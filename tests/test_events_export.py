"""Regression tests for GET /api/events/export.

Downloads whatever's already in the events table matching the same
selection fields sal/collectors/sm20.py's fetch_events() offers - a query
against local data, not a fresh RFC call. Uses the isolated_db fixture
(autouse, see conftest.py) with synthetic rows inserted directly, matching
tests/test_retention.py's _insert_event() convention.
"""
import io
from datetime import datetime, timezone

from openpyxl import load_workbook

from sal.storage.db import get_connection
from sal.web import create_app


def _client():
    app = create_app()
    app.testing = True
    return app.test_client()


def _insert_event(event_timestamp: str, counter: int = 1, source_system: str = "S23",
                   client: str = "100", user_id: str | None = None,
                   transaction_code: str | None = None, msg_code: str | None = None,
                   program: str | None = None, instance: str = "X",
                   message: str | None = None) -> None:
    conn = get_connection()
    try:
        conn.execute(
            "INSERT INTO events (source_system, client, instance, log_tstmp, counter, "
            "event_timestamp, user_id, transaction_code, msg_code, program, message, "
            "inserted_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (source_system, client, instance, event_timestamp, counter, event_timestamp,
             user_id, transaction_code, msg_code, program, message,
             datetime.now(timezone.utc).isoformat()),
        )
        conn.commit()
    finally:
        conn.close()


def test_events_export_requires_system_id():
    resp = _client().get("/api/events/export?dat_from=20260101&dat_to=20260101")
    assert resp.status_code == 400


def test_events_export_requires_valid_date_format():
    resp = _client().get("/api/events/export?system_id=S23&dat_from=2026-01-01&dat_to=20260101")
    assert resp.status_code == 400


def test_events_export_filters_by_date_range():
    _insert_event("2026-01-01 08:00:00", counter=1)  # outside range
    _insert_event("2026-01-05 08:00:00", counter=2)  # inside range
    _insert_event("2026-01-10 08:00:00", counter=3)  # outside range

    resp = _client().get(
        "/api/events/export?system_id=S23&dat_from=20260104&dat_to=20260106&format=csv"
    )
    assert resp.status_code == 200
    text = resp.get_data(as_text=True)
    assert "2026-01-05 08:00:00" in text
    assert "2026-01-01 08:00:00" not in text
    assert "2026-01-10 08:00:00" not in text


def test_events_export_filters_by_transaction_code():
    _insert_event("2026-01-05 08:00:00", counter=1, transaction_code="SU01")
    _insert_event("2026-01-05 08:00:01", counter=2, transaction_code="SM30")

    resp = _client().get(
        "/api/events/export?system_id=S23&dat_from=20260101&dat_to=20260131"
        "&transaction_code=SU01&format=csv"
    )
    text = resp.get_data(as_text=True)
    assert "SU01" in text
    assert "SM30" not in text


def test_events_export_filters_by_comma_separated_users():
    _insert_event("2026-01-05 08:00:00", counter=1, user_id="ALICE")
    _insert_event("2026-01-05 08:00:01", counter=2, user_id="BOB")
    _insert_event("2026-01-05 08:00:02", counter=3, user_id="CAROL")

    resp = _client().get(
        "/api/events/export?system_id=S23&dat_from=20260101&dat_to=20260131"
        "&user=ALICE, BOB&format=csv"
    )
    text = resp.get_data(as_text=True)
    assert "ALICE" in text
    assert "BOB" in text
    assert "CAROL" not in text


def test_events_export_scopes_by_client_and_system():
    _insert_event("2026-01-05 08:00:00", counter=1, source_system="S23", client="100")
    _insert_event("2026-01-05 08:00:01", counter=2, source_system="P01", client="100")
    _insert_event("2026-01-05 08:00:02", counter=3, source_system="S23", client="200")

    resp = _client().get(
        "/api/events/export?system_id=S23&client=100&dat_from=20260101&dat_to=20260131&format=csv"
    )
    text = resp.get_data(as_text=True)
    rows = [r for r in text.splitlines()[1:] if r.strip()]
    assert len(rows) == 1


def test_events_export_xlsx_returns_expected_headers_and_rows():
    _insert_event("2026-01-05 08:00:00", counter=1, transaction_code="SU01")

    resp = _client().get(
        "/api/events/export?system_id=S23&dat_from=20260101&dat_to=20260131&format=xlsx"
    )
    assert resp.status_code == 200
    wb = load_workbook(io.BytesIO(resp.data))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert header[0] == "event_timestamp"
    assert "transaction_code" in header
    assert ws.max_row == 2  # header + one event


def test_events_export_rejects_oversized_multi_value_filter():
    too_many = ",".join(f"U{i}" for i in range(501))
    resp = _client().get(
        f"/api/events/export?system_id=S23&dat_from=20260101&dat_to=20260131&user={too_many}"
    )
    assert resp.status_code == 400


def test_events_export_itgc_format_returns_cover_and_event_log_sheets():
    _insert_event("2026-01-05 08:00:00", counter=1, transaction_code="SU01", user_id="ALICE")

    resp = _client().get(
        "/api/events/export?system_id=S23&client=100&dat_from=20260101&dat_to=20260131"
        "&transaction_code=SU01&format=itgc"
    )
    assert resp.status_code == 200
    assert "sal_events_itgc.xlsx" in resp.headers["Content-Disposition"]
    wb = load_workbook(io.BytesIO(resp.data))
    assert wb.sheetnames == ["Cover Sheet", "Event Log"]

    cover = wb["Cover Sheet"]
    values = [row[2] for row in cover.iter_rows(min_row=1, max_col=3, values_only=True) if row[1]]
    assert any("SU01" in str(v) for v in values if v)  # extraction parameters shown

    log = wb["Event Log"]
    header = [c.value for c in log[1]]
    assert header[0] == "event_timestamp"
    assert log.max_row == 2  # header + one event


def test_events_export_defangs_formula_injection_in_message():
    _insert_event("2026-01-05 08:00:00", counter=1, message="=HYPERLINK(\"http://evil\")")

    resp = _client().get(
        "/api/events/export?system_id=S23&dat_from=20260101&dat_to=20260131&format=csv"
    )
    text = resp.get_data(as_text=True)
    assert "'=HYPERLINK" in text
