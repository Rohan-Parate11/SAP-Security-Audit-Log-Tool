"""Regression tests for sal/itgc_report.py.

Formatting is pure and doesn't touch the database, so these tests use
synthetic finding/collection-run dicts (matching test_findings.py's
_finding() convention) rather than the isolated_db fixture's sqlite file.
"""
import io

from openpyxl import load_workbook

import sal.itgc_report as itgc_report
from sal.rules import RULES


def _finding(rule_key="new_source", severity="High", user_id="U1",
             detected_at="2026-09-01T08:00:00+00:00", summary="synthetic finding",
             evidence=None, source_system="S23", client="100", status="open",
             first_seen_at="2026-09-01T08:00:00+00:00", last_seen_at="2026-09-01T08:00:00+00:00",
             reason_code=None, analyst_notes=None, disposed_by=None, disposed_at=None):
    return {
        "rule_key": rule_key,
        "rule_label": rule_key.replace("_", " ").title(),
        "severity": severity,
        "source_system": source_system,
        "client": client,
        "user_id": user_id,
        "detected_at": detected_at,
        "first_seen_at": first_seen_at,
        "last_seen_at": last_seen_at,
        "summary": summary,
        "evidence": evidence or {},
        "status": status,
        "reason_code": reason_code,
        "analyst_notes": analyst_notes,
        "disposed_by": disposed_by,
        "disposed_at": disposed_at,
    }


def _collection_run(run_id="R1", source_system="S23", client="100",
                     dat_from="20260801", dat_to="20260831",
                     started_at="2026-08-01T00:00:00+00:00",
                     finished_at="2026-08-01T00:05:00+00:00",
                     status="success", row_count=100, actor="tester",
                     filters_json=None, error_message=None):
    return {
        "run_id": run_id, "source_system": source_system, "client": client,
        "dat_from": dat_from, "dat_to": dat_to, "started_at": started_at,
        "finished_at": finished_at, "status": status, "row_count": row_count,
        "actor": actor, "filters_json": filters_json, "error_message": error_message,
    }


def test_control_map_covers_every_registered_rule():
    """Drift guard, mirroring sal/rules/_coverage.py's own drift test - a
    rule added to the registry without updating CONTROL_OBJECTIVES here
    would otherwise fall through to the silent UNMAPPED bucket.
    """
    mapped_keys = {key for objective in itgc_report.CONTROL_OBJECTIVES for key in objective["rules"]}
    registered_keys = {key for key, _label, _fn in RULES}
    assert registered_keys <= mapped_keys, f"unmapped rule keys: {registered_keys - mapped_keys}"


def test_control_map_refs_are_unique():
    refs = [objective["ref"] for objective in itgc_report.CONTROL_OBJECTIVES]
    assert len(refs) == len(set(refs))


def test_format_evidence_sorts_keys_and_rounds_floats():
    text = itgc_report._format_evidence({"ip": "10.0.0.1", "baseline_mean": 1.23456})
    assert text == "baseline_mean: 1.23; ip: 10.0.0.1"


def test_format_evidence_handles_empty():
    assert itgc_report._format_evidence(None) == ""
    assert itgc_report._format_evidence({}) == ""


def test_disposition_rationale_joins_reason_and_notes():
    f = _finding(reason_code="known_maintenance", analyst_notes="Cleared by SecOps")
    assert itgc_report._disposition_rationale(f) == "known_maintenance — Cleared by SecOps"


def test_disposition_rationale_handles_missing_parts():
    assert itgc_report._disposition_rationale(_finding()) == ""
    assert itgc_report._disposition_rationale(_finding(reason_code="x")) == "x"


def test_build_itgc_workbook_produces_expected_sheets_and_row_counts():
    findings = [
        _finding(rule_key="new_source", severity="High", status="open"),
        _finding(rule_key="mass_user_changes", severity="Medium", status="true_positive",
                  reason_code="confirmed", analyst_notes="Bulk SU01 change reviewed",
                  disposed_by="analyst1", disposed_at="2026-09-02T00:00:00+00:00"),
    ]
    collection_runs = [_collection_run()]

    data = itgc_report.build_itgc_workbook(
        findings, collection_runs, system_id="S23", client="100",
        environment="Production", generated_by="tester",
        generated_at="2026-09-11T00:00:00+00:00",
    )
    wb = load_workbook(io.BytesIO(data))

    assert wb.sheetnames == ["Cover Sheet", "Control Matrix", "Findings Detail", "Extraction Log"]

    detail = wb["Findings Detail"]
    assert detail.max_row == 1 + len(findings)  # header + one row per finding
    header = [c.value for c in detail[1]]
    assert header[:4] == ["S.No", "Control Ref", "ITGC Control Objective", "Test Performed (Rule)"]

    matrix = wb["Control Matrix"]
    assert matrix.max_row == 1 + len(itgc_report.CONTROL_OBJECTIVES)

    log = wb["Extraction Log"]
    assert log.max_row == 1 + len(collection_runs)


def test_control_matrix_population_counts_sum_correctly():
    findings = [
        _finding(rule_key="login_time", status="open"),
        _finding(rule_key="login_frequency", status="true_positive"),
        _finding(rule_key="new_source", status="false_positive"),
    ]
    data = itgc_report.build_itgc_workbook(
        findings, [], system_id="S23", client="100", generated_by="tester",
    )
    wb = load_workbook(io.BytesIO(data))
    matrix = wb["Control Matrix"]

    # Control 1.1 (Logon & Session Access Monitoring) covers all three rule keys above.
    row = next(r for r in matrix.iter_rows(min_row=2, values_only=True) if r[1] == "1.1")
    population, open_count, true_positive, false_positive = row[5], row[6], row[7], row[8]
    assert population == 3
    assert open_count == 1
    assert true_positive == 1
    assert false_positive == 1


def test_unmapped_rule_key_falls_back_to_other_bucket():
    findings = [_finding(rule_key="brand_new_rule_not_yet_mapped")]
    data = itgc_report.build_itgc_workbook(
        findings, [], system_id=None, client=None, generated_by="tester",
    )
    wb = load_workbook(io.BytesIO(data))
    detail = wb["Findings Detail"]
    row = [c.value for c in detail[2]]
    assert row[1] == itgc_report.UNMAPPED_REF
    assert row[2] == itgc_report.UNMAPPED_TITLE


def test_cover_sheet_testing_period_uses_successful_runs_only():
    """collection_runs.status is written as 'success'/'error' by
    sal/collectors/sm20.py's collect_and_store() - never 'completed'. This
    pins that against a real filter bug caught during live verification,
    where a 'completed' filter silently matched nothing.
    """
    collection_runs = [
        _collection_run(dat_from="20260801", dat_to="20260801", status="success"),
        _collection_run(dat_from="20260831", dat_to="20260831", status="success"),
        _collection_run(dat_from="20260915", dat_to="20260915", status="error"),
    ]
    data = itgc_report.build_itgc_workbook(
        [], collection_runs, system_id="S23", client="100", generated_by="tester",
    )
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Cover Sheet"]
    period_row = next(r for r in ws.iter_rows(values_only=True)
                       if r[1] == "Testing Period (SM20 data collected)")
    assert period_row[2] == "2026-08-01 to 2026-08-31"


def test_format_sap_date_converts_yyyymmdd_to_iso():
    assert itgc_report._format_sap_date("20260601") == "2026-06-01"


def test_format_sap_date_passes_through_unrecognized_values():
    assert itgc_report._format_sap_date("") == ""
    assert itgc_report._format_sap_date("not-a-date") == "not-a-date"


def test_findings_detail_defangs_formula_injection_from_disposition_notes():
    """Actor display names and disposition notes are unrestricted free text
    with no RBAC - and this workbook is built for handoff to the client's
    external audit team, so a formula like =HYPERLINK(...) reaching an
    opened cell there is a real risk, not just an internal-tool nuisance.
    """
    findings = [_finding(reason_code="=cmd|'/bin/calc'!A1", analyst_notes="ok",
                          disposed_by="=SUM(1,1)")]
    data = itgc_report.build_itgc_workbook(
        findings, [], system_id=None, client=None, generated_by="tester",
    )
    wb = load_workbook(io.BytesIO(data))
    detail = wb["Findings Detail"]
    row = [c.value for c in detail[2]]
    assert row[14].startswith("'=")  # Disposition Rationale
    assert row[15] == "'=SUM(1,1)"  # Reviewed By


def test_cover_sheet_defangs_formula_injection_from_generated_by():
    data = itgc_report.build_itgc_workbook(
        [], [], system_id=None, client=None, generated_by="=cmd|'/bin/calc'!A1",
    )
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Cover Sheet"]
    row = next(r for r in ws.iter_rows(values_only=True) if r[1] == "Report Generated By")
    assert row[2].startswith("'=")


def test_build_itgc_workbook_handles_no_findings_and_no_runs():
    data = itgc_report.build_itgc_workbook([], [], system_id=None, client=None)
    wb = load_workbook(io.BytesIO(data))
    assert wb["Findings Detail"].max_row == 1
    assert wb["Extraction Log"].max_row == 1
